import os

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

EXCEL_FILE_PATH = os.path.abspath(
    os.path.join(
        os.path.dirname(__file__),
        "..",
        "test_data",
        "GP_Derma_QA_Test_Cases.xlsx"
    )
)

def get_test_metadata(test_case_id):
    """
    Reads the Excel file and extracts metadata for the given test case ID.
    If the Excel file is unavailable or openpyxl is not installed, it falls back to a hardcoded dictionary.
    """
    if HAS_OPENPYXL and os.path.exists(EXCEL_FILE_PATH):
        try:
            wb = openpyxl.load_workbook(EXCEL_FILE_PATH, data_only=True)
            if "GP_Derma_QA" in wb.sheetnames:
                sheet = wb["GP_Derma_QA"]
            else:
                sheet = wb.active
                
            # Find column indices
            header_row = None
            headers = {}
            for i, row in enumerate(sheet.iter_rows(values_only=True)):
                if row and "Test Case ID" in row:
                    header_row = i
                    for col_idx, val in enumerate(row):
                        if val:
                            headers[str(val).strip()] = col_idx
                    break
            
            if header_row is not None:
                for row in sheet.iter_rows(min_row=header_row + 2, values_only=True):
                    row_id_idx = headers.get("Test Case ID")
                    if row_id_idx is not None and row[row_id_idx] == test_case_id:
                        return {
                            "test_case_id": test_case_id,
                            "module": row[headers.get("Module")] if "Module" in headers else "GP-Derma",
                            "feature": row[headers.get("Feature / Screen")] if "Feature / Screen" in headers else "",
                            "priority": row[headers.get("Priority")] if "Priority" in headers else "P2 - Medium",
                            "severity": row[headers.get("Severity")] if "Severity" in headers else "Medium",
                            "scenario": row[headers.get("Test Scenario / Objective")] if "Test Scenario / Objective" in headers else "",
                            "preconditions": row[headers.get("Pre-Conditions")] if "Pre-Conditions" in headers else "",
                            "steps": row[headers.get("Execution Steps")] if "Execution Steps" in headers else "",
                            "expected_result": row[headers.get("Expected Result")] if "Expected Result" in headers else "",
                            "tester_notes": row[headers.get("Tester Notes")] if "Tester Notes" in headers else "",
                        }
        except Exception as e:
            print(f"Failed to read from Excel: {e}. Falling back to default metadata.")
    
    # Fallback default metadata if Excel can't be read
    default_metadata = {
        "GPD-07": {
            "test_case_id": "GPD-07",
            "module": "GP-Derma",
            "feature": "Targeted Skin Lesion",
            "priority": "P0 - Blocker",
            "severity": "Critical",
            "scenario": "Verify AI clinical analysis for Targeted Skin Lesion",
            "preconditions": "User is logged in",
            "steps": "1. Select Anonymous patient\n2. Select Targeted skin lesion\n3. Upload image\n4. Enter details\n5. Analyze",
            "expected_result": "Treatment recommendations and/or barrier repair regimen should be populated.",
            "tester_notes": "Defect known: Treatment missing in some instances."
        },
        "GPD-08": {
            "test_case_id": "GPD-08",
            "module": "GP-Derma",
            "feature": "Melanoma Analysis",
            "priority": "P0 - Blocker",
            "severity": "Critical",
            "scenario": "Verify explicit biopsy verdict for melanoma",
            "preconditions": "User is logged in",
            "steps": "1. Upload melanoma lesion\n2. Analyze",
            "expected_result": "Explicit biopsy verdict must be shown.",
            "tester_notes": ""
        },
        # Add more defaults as needed
    }
    
    data = default_metadata.get(test_case_id, {})
    if not data:
        data = {
            "test_case_id": test_case_id,
            "module": "GP-Derma",
            "feature": "General",
            "priority": "P2 - Medium",
            "severity": "Medium",
            "scenario": "Execute test",
            "preconditions": "",
            "steps": "",
            "expected_result": "Test should pass.",
            "tester_notes": ""
        }
    return data
